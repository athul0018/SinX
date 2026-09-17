from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook

HEADER_MAP = {
    "poref": "po_ref",
    "equipmenttagno": "equipment_tag",
    "equipmentdescription": "description",
    "unit": "unit",
    "poquantity": "po_quantity",
    "completedquantity": "completed_quantity",
    "reminingquantity": "remaining_quantity",
    "remainingquantity": "remaining_quantity",
    "status": "status",
    "erectionfrontstatus": "erection_front_status",
    "remarks": "remarks",
    "photoref": "photo_ref",
}

TEMPLATE_HEADERS = [
    "PO ref",
    "Equipment Tag No.",
    "Equipment Description",
    "Unit",
    "P.O. Quantity",
    "Completed Quantity",
    "Remining Quantity",
    "Status",
    "Erection Front Status",
    "Remarks",
    "Photo Ref",
]


def _norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def _decimal(value) -> Decimal:
    if value is None or str(value).strip() == "":
        return Decimal("0")
    try:
        return Decimal(str(value).strip())
    except InvalidOperation:
        return Decimal("0")


def parse_master_excel(content: bytes) -> list[dict]:
    workbook = load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_norm(cell or "") for cell in rows[0]]
    mapped = [HEADER_MAP.get(header) for header in headers]
    items: list[dict] = []
    for raw in rows[1:]:
        if not raw or not any(raw):
            continue
        item = {
            "po_ref": "",
            "equipment_tag": "",
            "description": "",
            "unit": "",
            "po_quantity": Decimal("0"),
            "completed_quantity": Decimal("0"),
            "remaining_quantity": Decimal("0"),
            "status": "",
            "erection_front_status": "",
            "remarks": "",
            "photo_ref": "",
        }
        for index, key in enumerate(mapped):
            if not key or index >= len(raw):
                continue
            value = raw[index]
            if key in {"po_quantity", "completed_quantity", "remaining_quantity"}:
                item[key] = _decimal(value)
            else:
                item[key] = str(value or "").strip()
        if not item["po_ref"] and not item["equipment_tag"]:
            continue
        if item["remaining_quantity"] == 0 and item["po_quantity"]:
            remaining = item["po_quantity"] - item["completed_quantity"]
            item["remaining_quantity"] = remaining if remaining > 0 else Decimal("0")
        items.append(item)
    return items


def master_template_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Master"
    sheet.append(TEMPLATE_HEADERS)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
