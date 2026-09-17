from __future__ import annotations

import calendar
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

BLUE = PatternFill("solid", fgColor="BDD7EE")
HEADER_FONT = Font(bold=True, color="000000", name="Calibri", size=10)
BLACK = Font(name="Calibri", size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
VERTICAL = Alignment(horizontal="center", vertical="center", textRotation=90, wrap_text=True)
THIN = Border(
    left=Side(style="thin", color="9DC3E6"),
    right=Side(style="thin", color="9DC3E6"),
    top=Side(style="thin", color="9DC3E6"),
    bottom=Side(style="thin", color="9DC3E6"),
)
WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def _style_header(cell, *, vertical: bool = False) -> None:
    cell.fill = BLUE
    cell.font = HEADER_FONT
    cell.alignment = VERTICAL if vertical else CENTER
    cell.border = THIN


def attendance_workbook(
    site_name: str,
    year: int,
    month: int,
    employees: list[dict],
    marks: dict[tuple[str, int], dict],
) -> bytes:
    days = calendar.monthrange(year, month)[1]
    month_name = calendar.month_name[month]
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    # Layout: A=EMP.ID B=Name C=Designation, then each day uses 2 columns (status + OT)
    last_col = 3 + (days * 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = ws.cell(1, 1, f"{site_name} — Attendance Sheet — {month_name} {year}")
    title.font = Font(bold=True, size=14, name="Calibri")
    title.alignment = Alignment(horizontal="left", vertical="center")

    # Row 2: EMP.ID / Name / Designation + day numbers (merged across status+OT)
    for col, text in enumerate(["EMP.ID", "Employee Name", "Designation"], start=1):
        cell = ws.cell(2, col, text)
        _style_header(cell)
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)

    for day in range(1, days + 1):
        start = 3 + (day - 1) * 2 + 1
        end = start + 1
        ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=end)
        cell = ws.cell(2, start, day)
        _style_header(cell)
        ws.cell(2, end).fill = BLUE
        ws.cell(2, end).border = THIN

        weekday = WEEKDAYS[date(year, month, day).weekday()]
        day_cell = ws.cell(3, start, weekday)
        ot_cell = ws.cell(3, end, "OT")
        _style_header(day_cell, vertical=True)
        _style_header(ot_cell, vertical=True)
        ws.column_dimensions[get_column_letter(start)].width = 4
        ws.column_dimensions[get_column_letter(end)].width = 4

    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 45
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16

    for row_index, emp in enumerate(employees, start=4):
        for col, value in enumerate([emp["employee_code"], emp["name"], emp["designation"]], start=1):
            cell = ws.cell(row_index, col, value)
            cell.font = BLACK
            cell.border = THIN
            cell.alignment = Alignment(vertical="center")
        for day in range(1, days + 1):
            rec = marks.get((emp["id"], day)) or {}
            status_col = 3 + (day - 1) * 2 + 1
            ot_col = status_col + 1
            status_cell = ws.cell(row_index, status_col, rec.get("mark") or "")
            ot_value = rec.get("ot")
            ot_cell = ws.cell(
                row_index,
                ot_col,
                ot_value if ot_value not in (None, "", 0, 0.0) else "",
            )
            for cell in (status_cell, ot_cell):
                cell.alignment = CENTER
                cell.border = THIN
                cell.font = BLACK

    legend_row = 5 + len(employees)
    ws.cell(
        legend_row,
        1,
        "F = Full day    H = Half day    A = Absent    blank = not marked    OT = overtime hours",
    )

    return _save(wb)


def progress_workbook(
    site_name: str,
    year: int,
    month: int,
    progress_rows: list[list],
    master_rows: list[list],
    non_po_rows: list[list],
) -> bytes:
    month_name = calendar.month_name[month]
    wb = Workbook()

    _fill_sheet(
        wb.active,
        "Progress",
        f"{site_name} — Monthly Progress — {month_name} {year}",
        [
            "Date",
            "Plan ID",
            "Classification",
            "PO Ref",
            "Equipment Tag No.",
            "Job Description",
            "Status",
            "Work Front Status",
            "Quantity Completed",
            "Hours",
            "Remarks",
            "Photo Ref",
            "Reported By",
        ],
        progress_rows,
    )
    master_ws = wb.create_sheet("PO Master")
    _fill_sheet(
        master_ws,
        "PO Master",
        f"{site_name} — PO Master List",
        [
            "PO ref",
            "Equipment Tag No.",
            "Equipment Description",
            "Unit",
            "P.O. Quantity",
            "Completed Quantity",
            "Remining Quantity",
            "Status",
            "Erection Front Status",
            "Remarks",
            "Photo Ref",
        ],
        master_rows,
    )
    non_po_ws = wb.create_sheet("Non-PO Jobs")
    _fill_sheet(
        non_po_ws,
        "Non-PO Jobs",
        f"{site_name} — Non-PO jobs — {month_name} {year}",
        [
            "Date",
            "Clarification",
            "Job Description",
            "Allocated Workers",
            "Status",
            "Remarks",
            "Photo Ref",
        ],
        non_po_rows,
    )
    return _save(wb)


def _fill_sheet(ws: Worksheet, title: str, heading: str, headers: list[str], rows: list[list]) -> None:
    ws.title = title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    cell = ws.cell(1, 1, heading)
    cell.font = Font(bold=True, size=14, name="Calibri")
    for col, text in enumerate(headers, start=1):
        header = ws.cell(2, col, text)
        _style_header(header)
        ws.column_dimensions[get_column_letter(col)].width = max(14, min(28, len(text) + 4))
    for r_index, row in enumerate(rows, start=3):
        for c_index, value in enumerate(row, start=1):
            cell = ws.cell(r_index, c_index, value)
            cell.font = BLACK
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def _save(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
